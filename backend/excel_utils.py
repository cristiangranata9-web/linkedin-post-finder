"""
Lettura del file .xlsx di input (colonna "Email") e generazione dei due file
.xlsx di output (vista settimanale e vista dettagliata), con link cliccabili.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

GIORNI_IT = ["Lunedì", "Martedì", "Mercoledì", "Giovedì", "Venerdì", "Sabato", "Domenica"]


class InputFileError(ValueError):
    pass


def read_emails_from_xlsx(file_bytes: bytes) -> list[str]:
    """Legge la colonna 'Email' dal primo foglio del file caricato."""
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise InputFileError(f"Impossibile leggere il file .xlsx: {exc}") from exc

    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise InputFileError("Il file è vuoto.")

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    try:
        email_col = next(i for i, h in enumerate(header) if h.lower() == "email")
    except StopIteration as exc:
        raise InputFileError(
            "Non trovo una colonna intestata 'Email' nel file caricato. "
            f"Intestazioni trovate: {header}"
        ) from exc

    emails: list[str] = []
    for row in rows[1:]:
        if email_col >= len(row):
            continue
        value = row[email_col]
        if value is None:
            continue
        value = str(value).strip()
        if value:
            emails.append(value)

    if not emails:
        raise InputFileError("La colonna 'Email' non contiene indirizzi validi.")

    return emails


HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="1D4ED8", underline="single")


def _style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def build_weekly_xlsx(weekly_rows: list[dict[str, Any]]) -> bytes:
    """
    Colonne: Email | Nome profilo LinkedIn | URL profilo LinkedIn |
             Lunedì..Domenica (link multipli impilati nella stessa cella).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Vista settimanale"

    headers = ["Email", "Nome profilo LinkedIn", "URL profilo LinkedIn"] + GIORNI_IT
    ws.append(headers)
    _style_header(ws, len(headers))

    for row in weekly_rows:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=row.get("email", ""))
        ws.cell(row=r, column=2, value=row.get("profile_name") or row.get("status_label", ""))
        url = row.get("profile_url")
        if url:
            c = ws.cell(row=r, column=3, value=url)
            c.hyperlink = url
            c.font = LINK_FONT
        else:
            ws.cell(row=r, column=3, value=row.get("status_label", ""))

        giorni = row.get("giorni", {})
        for idx, giorno in enumerate(GIORNI_IT):
            links = giorni.get(giorno, [])
            col = 4 + idx
            cell = ws.cell(row=r, column=col)
            if links:
                cell.value = "\n".join(links)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                # openpyxl non supporta più hyperlink diversi nella stessa cella:
                # il testo elenca gli URL completi (cliccabili una volta incollati/
                # aperti dal browser) uno per riga; il primo link viene reso cliccabile.
                cell.hyperlink = links[0]
                cell.font = LINK_FONT
            else:
                cell.value = ""

    widths = [28, 24, 40] + [30] * 7
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_detailed_xlsx(detailed_rows: list[dict[str, Any]]) -> bytes:
    """
    Colonne: Email | Nome profilo LinkedIn | URL profilo LinkedIn |
             Data del post | Giorno della settimana | Link al post
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Vista dettagliata"

    headers = [
        "Email",
        "Nome profilo LinkedIn",
        "URL profilo LinkedIn",
        "Data del post",
        "Giorno della settimana",
        "Link al post",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for row in detailed_rows:
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=row.get("email", ""))
        ws.cell(row=r, column=2, value=row.get("profile_name", ""))
        url = row.get("profile_url")
        if url:
            c = ws.cell(row=r, column=3, value=url)
            c.hyperlink = url
            c.font = LINK_FONT
        else:
            ws.cell(row=r, column=3, value="")
        ws.cell(row=r, column=4, value=row.get("post_date", ""))
        ws.cell(row=r, column=5, value=row.get("day_of_week", ""))
        post_link = row.get("post_link")
        if post_link:
            c = ws.cell(row=r, column=6, value=post_link)
            c.hyperlink = post_link
            c.font = LINK_FONT
        else:
            ws.cell(row=r, column=6, value="")

    widths = [28, 24, 40, 16, 20, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
